from io import BytesIO
from docx import Document
import fitz

class ResumeParseError(ValueError): pass

def parse_resume(filename: str, content: bytes) -> str:
    name = filename.lower()
    try:
        if name.endswith(".pdf"):
            with fitz.open(stream=content, filetype="pdf") as pdf:
                text = "\n".join(page.get_text("text") for page in pdf)
        elif name.endswith(".docx"):
            doc = Document(BytesIO(content))
            text = "\n".join(p.text for p in doc.paragraphs)
            for table in doc.tables:
                for row in table.rows:
                    text += "\n" + " | ".join(cell.text for cell in row.cells)
        else: raise ResumeParseError("Only PDF and DOCX resumes are supported.")
    except ResumeParseError: raise
    except Exception as exc: raise ResumeParseError("We could not read this resume. Please try another PDF or DOCX file.") from exc
    text = "\n".join(line.strip() for line in text.splitlines() if line.strip())
    if len(text) < 30: raise ResumeParseError("Very little text was extracted. If this is a scanned PDF, use a text-based PDF or DOCX version.")
    return text[:50000]
